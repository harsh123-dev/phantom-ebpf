from docx import Document
p=r'D:\project-1\tmp\phantom-template\pandoc-draft.docx'; d=Document(p)
for i,p in enumerate(d.paragraphs):
 if i>=78 and p.text.strip(): print(i,p.style.name,repr(p.text[:180]))
for i,t in enumerate(d.tables):
 print('TABLE',i,'rows',len(t.rows),'cols',len(t.columns),'first', [c.text[:30] for c in t.rows[0].cells])
