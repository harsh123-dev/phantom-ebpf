from docx import Document
from zipfile import ZipFile
from pathlib import Path
for label,path in [('template',r'D:\Downloads\BCSE497J Project I Report - Template FINAlL Updated  (2).docx'),('draft',r'D:\project-1\tmp\phantom-template\pandoc-draft.docx')]:
 d=Document(path); print(label,'sections',len(d.sections))
 for i,s in enumerate(d.sections):
  print(i,[round(getattr(s,x)/914400,3) for x in ('page_width','page_height','left_margin','right_margin','top_margin','bottom_margin')], 'header', [p.text for p in s.header.paragraphs], 'footer',[p.text for p in s.footer.paragraphs], 'different',s.different_first_page_header_footer)
 with ZipFile(path) as z:
  print('parts', [x for x in z.namelist() if x.startswith('word/header') or x.startswith('word/footer') or 'media/' in x])
  print('fields', {k:z.read(k).decode('utf-8','ignore').count('PAGE') for k in z.namelist() if k.endswith('.xml') and (k.startswith('word/header') or k.startswith('word/footer') or k=='word/document.xml')})
