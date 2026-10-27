from docx import Document
p=r'D:\project-1\tmp\phantom-template\pandoc-draft.docx'
d=Document(p)
print('paras',len(d.paragraphs),'tables',len(d.tables),'sections',len(d.sections))
for i,p in enumerate(d.paragraphs):
 if p.text.strip(): print(i,p.style.name,repr(p.text[:140]))
