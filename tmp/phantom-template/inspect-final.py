from docx import Document
p=r'D:\project-1\output\PHANTOM_Project_Report_Word\PHANTOM_Project_Report.docx';d=Document(p)
print('paras',len(d.paragraphs),'tables',len(d.tables),'sections',len(d.sections))
for i,p in enumerate(d.paragraphs[:35]): print(i,p.style.name,repr(p.text[:100]))
for i,t in enumerate(d.tables): print('TABLE',i,len(t.rows),len(t.columns),[c.text[:40] for c in t.rows[0].cells] if t.rows else [])
