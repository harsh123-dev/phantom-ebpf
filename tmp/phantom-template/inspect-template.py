from docx import Document
from pathlib import Path
p=Path(r'D:\Downloads\BCSE497J Project I Report - Template FINAlL Updated  (2).docx')
d=Document(str(p))
print('paragraphs',len(d.paragraphs),'tables',len(d.tables),'sections',len(d.sections))
for i,s in enumerate(d.sections):
 print('SECTION',i,'size',s.page_width/914400,s.page_height/914400,'margins',*(round(getattr(s,x)/914400,3) for x in ['left_margin','right_margin','top_margin','bottom_margin']),'break',s.start_type)
for i,p in enumerate(d.paragraphs):
 t=p.text.replace('\n',' / ')
 if t.strip(): print('P',i,repr(t[:200]),'STYLE',p.style.name)
for i,t in enumerate(d.tables):
 print('TABLE',i,len(t.rows),len(t.columns))
 for row in t.rows[:6]: print(' | '.join(c.text[:80].replace('\n','/') for c in row.cells))
